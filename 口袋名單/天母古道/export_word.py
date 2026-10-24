"""將同目錄的規劃筆記匯出為 Word。使用專案 .venv 執行。"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def add_text(paragraph, text):
    """保留 Markdown 連結為可點選的 Word 超連結。"""
    start = 0
    for match in re.finditer(r"\[([^\]]+)\]\(([^)]+)\)", text):
        paragraph.add_run(text[start:match.start()])
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        props.append(color)
        run.append(props)
        node = OxmlElement("w:t")
        node.text = match[1]
        run.append(node)
        link.append(run)
        paragraph._p.append(link)
        start = match.end()
    paragraph.add_run(text[start:])


base = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(1.7)
for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
    style = doc.styles[name]
    style.font.name = "Microsoft JhengHei"
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
doc.styles["Normal"].font.size = Pt(10)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].font.size = Pt(14)
doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("245C42")

table = None
for line in (base / "README.md").read_text().splitlines():
    if line.startswith("| "):
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            continue
        if table is None:
            table = doc.add_table(rows=0, cols=len(cells))
            table.style = "Table Grid"
        row = table.add_row()
        for cell, text in zip(row.cells, cells):
            cell.text = text
        if len(table.rows) == 1:
            for cell in row.cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        continue
    table = None
    if not line:
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:], level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=2)
    elif line.startswith("- "):
        add_text(doc.add_paragraph(style="List Bullet"), line[2:])
    else:
        add_text(doc.add_paragraph(), line)

output = base / "天母古道_口袋名單.docx"
doc.save(output)
# 確認產出檔案可重新解析，且表格與正文均完整。
check = Document(output)
assert len(check.tables) == 1 and len(check.tables[0].rows) == 3
assert any("目前規劃" in p.text for p in check.paragraphs)
print(f"已建立：{output}\n檔案大小：{output.stat().st_size:,} bytes")
