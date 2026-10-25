from pathlib import Path
import re
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

base = Path(__file__).resolve().parent
root = base.parents[1]
source = base / "README.md"
out = base / "2026-11-01_坪頂古圳健行.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.8)

for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "List Number"):
    style = doc.styles[name]
    style.font.name = "Microsoft JhengHei"
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
doc.styles["Normal"].font.size = Pt(10)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].font.color.rgb = RGBColor(36, 92, 66)


def add_inline(paragraph, text):
    pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|(\*\*.+?\*\*)")
    position = 0
    for match in pattern.finditer(text):
        paragraph.add_run(text[position:match.start()])
        if match.group(1):
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(match.group(2), RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "0563C1")
            props.append(color)
            run.append(props)
            value = OxmlElement("w:t")
            value.text = match.group(1)
            run.append(value)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            paragraph.add_run(match.group(3).strip("*")).bold = True
        position = match.end()
    paragraph.add_run(text[position:])


def table_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line):
    cells = table_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(lines):
    header = table_cells(lines[0])
    rows = [table_cells(line) for line in lines[2:]]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for column, value in enumerate(header):
        cell = header_row.cells[column]
        add_inline(cell.paragraphs[0], value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E2EEE7")
        cell._tc.get_or_add_tcPr().append(shade)
    for values in rows:
        cells = table.add_row().cells
        for column, value in enumerate(values[:len(header)]):
            add_inline(cells[column].paragraphs[0], value)
    return table


lines = source.read_text(encoding="utf-8").splitlines()
index = 0
while index < len(lines):
    line = lines[index]
    if not line.strip():
        index += 1
        continue
    if line.strip().startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
        table_lines = [line, lines[index + 1]]
        index += 2
        while index < len(lines) and lines[index].strip().startswith("|"):
            table_lines.append(lines[index])
            index += 1
        add_table(table_lines)
        continue
    if "<!-- WORD_PRIVATE_CHILD_AGE -->" in line:
        age_script = root / ".opencode/skills/family-hiking/scripts/child_age.py"
        python = root / ".venv/bin/python"
        private_notes = root / ".private/共用筆記.md"
        if private_notes.exists():
            if not python.exists():
                python = Path(sys.executable)
            result = subprocess.run(
                [str(python), str(age_script), base.name],
                check=True, capture_output=True, text=True,
            )
            age_text = result.stdout.strip().split("：", 1)[-1]
            age_line = f"幼兒活動日年齡（本機私密 Word，{base.name}）：{age_text}"
        else:
            age_line = "幼兒年齡與背負裝備適用性：活動日前依同行者私密資料及裝備規範確認。"
        paragraph = doc.add_paragraph()
        paragraph.add_run(age_line).bold = True
        index += 1
        continue
    image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
    if image:
        image_path = base / image.group(2)
        if image_path.suffix.lower() == ".svg":
            converter = shutil.which("rsvg-convert")
            if not converter:
                raise SystemExit("匯入 SVG 路線圖需要安裝 rsvg-convert，才能轉成 Word 可用圖片。")
            rendered_dir = root / ".cache" / "word"
            rendered_dir.mkdir(parents=True, exist_ok=True)
            rendered = rendered_dir / f"{base.name}-route.png"
            subprocess.run([converter, "-w", "1400", "-o", str(rendered), str(image_path)], check=True)
            image_path = rendered
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Cm(16))
    elif line.startswith("# "):
        doc.add_heading(line[2:], 0)
    elif line.startswith("## "):
        doc.add_heading(line[3:], 1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], 2)
    elif re.match(r"^\d+\. ", line):
        paragraph = doc.add_paragraph(style="List Number")
        add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
    elif line.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline(paragraph, line[2:])
    else:
        add_inline(doc.add_paragraph(), line.removeprefix("> "))
    index += 1

doc.save(out)
check = Document(out)
assert any("坪頂古圳" in p.text for p in check.paragraphs)
assert len(check.inline_shapes) == 4
assert len(check.tables) == 1 and len(check.tables[0].rows) == 5
assert any("幼兒活動日年齡（本機私密 Word" in p.text for p in check.paragraphs)
assert any(rel.reltype == RT.HYPERLINK and "maps.app.goo.gl" in rel.target_ref for rel in check.part.rels.values())
assert any(rel.reltype == RT.HYPERLINK and "ymsnp.gov.tw" in rel.target_ref for rel in check.part.rels.values())
assert any(rel.reltype == RT.HYPERLINK and "坪頂古圳_等高線圖.pdf" in rel.target_ref for rel in check.part.rels.values())
assert any(rel.reltype == RT.HYPERLINK and "坪頂古圳_山頂方位圖.pdf" in rel.target_ref for rel in check.part.rels.values())
print(f"已產生並驗證：{out} ({out.stat().st_size:,} bytes)")
