from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

base = Path(__file__).resolve().parent
source = base / 'README.md'
out = base / '2026-10-25_大溝溪環線規劃.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.8)
for name in ('Normal','Title','Heading 1','Heading 2','List Bullet','List Number'):
    st=doc.styles[name]
    st.font.name='Microsoft JhengHei'
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].font.color.rgb=RGBColor(36,92,66)

def inline(p,s):
    # Simple bold and markdown links, preserving link clickability.
    token=re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*(.+?)\*\*')
    pos=0
    for m in token.finditer(s):
        p.add_run(s[pos:m.start()])
        if m.group(1):
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),p.part.relate_to(m.group(2),RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rp.append(c); r.append(rp)
            t=OxmlElement('w:t'); t.text=m.group(1); r.append(t); link.append(r); p._p.append(link)
        else:
            p.add_run(m.group(3)).bold=True
        pos=m.end()
    p.add_run(s[pos:])

lines=source.read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells): rows.append(cells)
            i+=1
        if rows:
            t=doc.add_table(rows=0,cols=len(rows[0])); t.style='Table Grid'
            for ri,rowdata in enumerate(rows):
                cells=t.add_row().cells
                for cell,text in zip(cells,rowdata):
                    cell.text=text
                    if ri==0:
                        for r in cell.paragraphs[0].runs:r.bold=True
                if ri==0:
                    trPr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
        continue
    if line.startswith('# '): doc.add_heading(line[2:],0)
    elif line.startswith('## '): doc.add_heading(line[3:],1)
    elif line.startswith('### '): doc.add_heading(line[4:],2)
    elif line.startswith('> '):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.4); inline(p,line[2:])
    elif re.match(r'^\d+\. ',line):
        p=doc.add_paragraph(style='List Number'); inline(p,re.sub(r'^\d+\. ','',line))
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:])
    else: inline(doc.add_paragraph(),line)
    i+=1

doc.save(out)
check=Document(out)
assert len(check.tables)==1 and len(check.tables[0].rows)==5
assert any('出發時間 ABCD' in p.text for p in check.paragraphs)
print(f'已產生並驗證：{out} ({out.stat().st_size:,} bytes)')
